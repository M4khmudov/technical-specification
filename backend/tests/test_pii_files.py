"""B2, S2: персональные данные не уходят в файлы для Code Interpreter и в логи;
псевдонимизация не портит данные и не путает людей."""

import io
import json
import logging
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from app.assistant.code_interpreter import mask_file, validate_upload
from app.assistant.pii import Pseudonymizer
from app.core.errors import DomainError
from app.core.logging import StripQueryFilter

PII = Pseudonymizer(
    {1: "Иванова Мария Петровна", 2: "Петров Иван Сергеевич", 3: "Кузнецова Анна Игоревна"},
    {
        1: "ООО «Восковой двор»",
        2: "ООО «Стеклотара Плюс»",
        3: "ИП Смирнова Ольга Викторовна",
        4: "Global Wicks Ltd",
        5: "Aroma Trade GmbH",
    },
)
NAMES = ["Петров", "Иванова", "Кузнецова", "Смирнова", "Восковой", "Стеклотара", "Wicks"]


def _all_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", "replace") for name in archive.namelist()
        )


def _leaks(text: str) -> list[str]:
    return [n for n in NAMES if n.lower() in text.lower()]


# ---------------------------------------------------------------- XLSX: всё, где бывает текст


def _dirty_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Петров"
    ws["A1"] = "Ответственный: Петров Иван Сергеевич"
    ws["A2"] = '="Кузнецова"&" А."'
    ws["A3"] = 42
    ws["A3"].comment = Comment("Проверила Иванова Мария", "Иванова Мария Петровна")
    ws.oddHeader.center.text = "Отчёт: Смирнова Ольга Викторовна"
    ws.oddFooter.left.text = "ООО «Восковой двор»"
    wb.properties.creator = "Петров Иван Сергеевич"
    wb.properties.lastModifiedBy = "Кузнецова Анна Игоревна"
    wb.properties.title = "Поставки Global Wicks Ltd"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_xlsx_has_no_names_anywhere_inside():
    masked = mask_file("report.xlsx", _dirty_workbook(), PII)
    assert _leaks(_all_xml(masked)) == []
    # данные остались: значения и числа на месте, имена — псевдонимами
    ws = load_workbook(io.BytesIO(masked)).active
    assert ws["A1"].value == "Ответственный: Сотрудник-2"
    assert ws["A3"].value == 42


def test_formula_text_is_masked_when_there_is_no_cached_value():
    masked = mask_file("f.xlsx", _dirty_workbook(), PII)
    ws = load_workbook(io.BytesIO(masked)).active
    assert "Кузнецова" not in str(ws["A2"].value)


# ---------------------------------------------------------------- CSV в разных кодировках


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251", "utf-16", "utf-16-le"])
def test_csv_in_any_encoding_is_masked(encoding):
    data = "Кто;Сколько\nПетров Иван Сергеевич;10\nООО «Восковой двор»;5\n".encode(encoding)
    masked = mask_file("data.csv", data, PII).decode("utf-8")
    assert _leaks(masked) == []
    assert "Сотрудник-2" in masked and "Поставщик-1" in masked


def test_unreadable_csv_is_rejected_not_passed_through():
    data = bytes(range(256)) * 4  # не текст ни в одной кодировке
    with pytest.raises(DomainError):
        mask_file("bin.csv", b"\x00\xff" + data, PII)


# ---------------------------------------------------------------- склеенные написания


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Петров_отчёт.xlsx", "Сотрудник-2_отчёт.xlsx"),
        ("Кузнецова_А.csv", "Сотрудник-3_А.csv"),
        ("Восковой_двор_прайс.csv", "Поставщик-1_прайс.csv"),
        ("ВосковойДвор", "Поставщик-1"),
        ("Petrov_report.csv", "Сотрудник-2_report.csv"),
        ("Отгрузка от Voskovoy dvor", "Отгрузка от Поставщик-1"),
    ],
)
def test_glued_and_translit_names_are_masked(text, expected):
    assert PII.mask(text) == expected


# ---------------------------------------------------------------- S2: не портить данные


@pytest.mark.parametrize(
    "supplier", ["Global Wicks Ltd", "Aroma Trade GmbH", "ООО «Восковой двор»"]
)
def test_masking_keeps_json_valid(supplier):
    text = json.dumps({"supplier": supplier, "note": f"от {supplier}"}, ensure_ascii=False)
    masked = PII.mask(text)
    json.loads(masked)  # не падает: кавычки не «съедены»


def test_company_name_is_not_taken_for_a_person():
    """«ООО «Стеклотара Плюс»» — не ФИО: слово «стеклотара» в вопросе остаётся словом."""
    assert PII.mask("Сколько у нас стеклотары?") == "Сколько у нас стеклотары?"
    assert PII.mask("ООО «Стеклотара Плюс»") == "Поставщик-2"


def test_namesake_is_not_turned_into_an_employee():
    """«Петрова Ольга Николаевна» — не Петров Иван Сергеевич."""
    text = "Принимала Петрова Ольга Николаевна"
    assert PII.unmask(PII.mask(text)) == text


# ---------------------------------------------------------------- B2: логи и загрузка


def test_access_log_has_no_query_string():
    record = logging.LogRecord(
        "uvicorn.access",
        logging.INFO,
        "",
        0,
        '%s - "%s %s HTTP/%s" %d',
        ("1.2.3.4:5", "GET", "/api/batches?supplier=Смирнова&page=1", "1.1", 200),
        None,
    )
    assert StripQueryFilter().filter(record)
    assert "Смирнова" not in record.getMessage()
    assert "/api/batches" in record.getMessage()


def test_xlsx_bomb_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 60_000_000)
    with pytest.raises(DomainError, match="слишком"):
        validate_upload("bomb.xlsx", buffer.getvalue())
