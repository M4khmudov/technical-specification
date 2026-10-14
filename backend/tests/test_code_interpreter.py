"""Code Interpreter: данные в контейнер с псевдонимами, файлы модели обратно к нам,
показ кода, ошибка кода не роняет чат."""

import io
import json
from datetime import date, timedelta
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from app.assistant.code_interpreter import CodeInterpreterSession, mask_file
from app.assistant.pii import Pseudonymizer
from app.assistant.tools import ToolContext, default_registry
from app.core import clock
from app.core.errors import DomainError
from app.llm.fake import ScriptedProvider
from app.llm.registry import get_llm
from app.llm.router import LLMRouter, Route
from app.llm.types import (
    CodeExecuted,
    Completed,
    FileCited,
    TextDelta,
    ToolCall,
    ToolCallRequested,
)
from app.main import app
from app.models import ChatMessage, Conversation, Supplier, User
from app.models.enums import Currency, MessageRole, Role
from app.schemas.production import ActualIn, JobCompleteIn, JobCreateIn
from app.schemas.stock import ReceiptIn, ReceiptLineIn
from app.services import batches, exports, production
from tests.factories import make_material, make_recipe, make_supplier, make_user
from tests.fakes import FakeRates
from tests.test_chat import parse_sse

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 32
PII = Pseudonymizer({7: "Лебедев Олег Ильич"}, {9: "ООО «Медовый луг»"})


# ---------------------------------------------------------------- наборы данных


async def _receipt(session, user, supplier, material, qty="10", price="100", day=None):
    receipt = await batches.create_receipt(
        session,
        user,
        ReceiptIn(
            supplier_id=supplier.id,
            received_on=day or date(2026, 9, 1),
            lines=[
                ReceiptLineIn(
                    material_id=material.id,
                    quantity=Decimal(qty),
                    price=Decimal(price),
                    currency=Currency.RUB,
                )
            ],
        ),
        FakeRates(),
    )
    for b in receipt.batches:
        await batches.accept_batch(session, user, b.id)
    return receipt


async def test_consumption_dataset_matches_completed_job(session):
    manager = await make_user(session, role=Role.MANAGER)
    supplier = await make_supplier(session)
    wax = await make_material(session, code=f"CI-{uuid4().hex[:5]}", name="Воск ci")
    recipe = await make_recipe(session, {wax: "0.5"})
    await _receipt(session, manager, supplier, wax, qty="20", price="300")
    job = await production.create_job(
        session,
        manager,
        JobCreateIn(recipe_id=recipe.id, quantity=10, planned_on=clock.today()),
        plan=True,
    )
    await production.start_job(session, manager, job.id)
    await production.complete_job(
        session,
        manager,
        job.id,
        JobCompleteIn(actuals=[ActualIn(material_id=wax.id, actual_qty=6)]),
    )

    table = await exports.consumption(session, None, None, wax.code)

    assert table.columns[:6] == [
        "date",
        "week_start",
        "material_code",
        "material_name",
        "unit",
        "quantity",
    ]
    # 5 кг из резерва + 1 кг перерасхода — две строки списания
    assert [r[5] for r in table.rows] == [Decimal("5.000"), Decimal("1.000")]
    assert sum(r[7] for r in table.rows) == Decimal("1800.00")
    row = table.rows[0]
    assert (row[2], row[8]) == (wax.code, job.number)
    assert row[1] == row[0] - timedelta(days=row[0].weekday())  # понедельник недели
    assert f"{wax.code},Воск ci,кг,5," in table.to_csv()  # числа с точкой, без экспоненты


async def test_recipe_costs_use_average_free_stock_cost(session):
    manager = await make_user(session)
    supplier = await make_supplier(session)
    wax = await make_material(session, code=f"CI-{uuid4().hex[:5]}")
    recipe = await make_recipe(session, {wax: "0.2"}, name=f"Рецепт ci {uuid4().hex[:4]}")
    await _receipt(session, manager, supplier, wax, qty="10", price="100")
    await _receipt(session, manager, supplier, wax, qty="30", price="200")

    table = await exports.recipe_costs(session)

    row = next(r for r in table.rows if r[0] == recipe.name)
    # (10×100 + 30×200) / 40 = 175 ₽/кг; на свечу 0,2 кг → 35 ₽
    assert row[5] == Decimal("175.000000")
    assert row[6] == Decimal("35.0000")


async def test_too_long_period_is_rejected(session):
    with pytest.raises(DomainError, match="слишком длинный"):
        await exports.consumption(session, date(2020, 1, 1), date(2026, 1, 1), None)


# ---------------------------------------------------------------- маскирование файлов


def test_csv_for_container_has_pseudonyms():
    data = "поставщик;принял\nООО «Медовый луг»;Лебедев О. И.\n".encode("cp1251")
    masked = mask_file("Отчёт.csv", data, PII).decode("utf-8")
    assert masked == "поставщик;принял\nПоставщик-9;Сотрудник-7\n"


def test_xlsx_for_container_has_pseudonyms_in_every_cell():
    wb = Workbook()
    ws = wb.active
    ws.title = "Медовый луг"
    ws.append(["Поставщик", "Кг", "Принял"])
    ws.append(["Медовый луг, ООО", 12.5, "Лебедеву О. И."])
    ws.append(["=SUM(B2:B2)", None, None])
    buffer = io.BytesIO()
    wb.save(buffer)

    masked = load_workbook(io.BytesIO(mask_file("x.xlsx", buffer.getvalue(), PII)))
    sheet = masked.active
    assert sheet.title == "Поставщик-9"
    assert [c.value for c in sheet[2]] == ["Поставщик-9", 12.5, "Сотрудник-7"]
    assert sheet["A3"].value == "=SUM(B2:B2)"  # формулы не трогаем


async def test_export_tool_uploads_masked_file(session):
    manager = await make_user(session, role=Role.STOREKEEPER, full_name="Лебедев Олег Ильич")
    supplier = Supplier(name=f"ООО «Медовый луг {uuid4().hex[:4]}»")
    session.add(supplier)
    await session.flush()
    wax = await make_material(session, code=f"CI-{uuid4().hex[:5]}")
    await _receipt(session, manager, supplier, wax, day=clock.today() - timedelta(days=3))
    provider = ScriptedProvider()
    pii = Pseudonymizer({manager.id: manager.full_name}, {supplier.id: supplier.name})
    ci = CodeInterpreterSession(provider, pii)
    conversation = Conversation(user_id=manager.id, title="t")
    session.add(conversation)
    await session.flush()
    message = ChatMessage(conversation_id=conversation.id, role=MessageRole.ASSISTANT)
    session.add(message)
    await session.flush()
    ctx = ToolContext(session, manager, conversation.id, message.id, ci=ci)

    outcome, _ = await default_registry().execute(
        ToolCall("c1", "export_data", json.dumps({"dataset": "receipts", "material": wax.code})),
        ctx,
    )

    assert outcome.ok, outcome.output
    assert outcome.output["rows"] == 1
    # реальный путь в контейнере OpenAI: /mnt/data/<file_id>-<имя>
    assert outcome.output["file"].startswith("/mnt/data/file-1-receipts_")
    # точные коды материалов: модель фильтрует по коду, а не по подстроке названия
    assert outcome.output["materials"] == {wax.code: wax.name}
    ((filename, content),) = provider.uploaded.values()
    text = content.decode()
    assert "Медовый луг" not in text and f"Поставщик-{supplier.id}" in text
    assert ci.spec().file_ids == ("file-1",)


async def test_export_tool_without_code_interpreter(session):
    user = await make_user(session)
    ctx = ToolContext(session, user, 0, 0, ci=None)
    outcome, _ = await default_registry().execute(
        ToolCall("c1", "export_data", '{"dataset": "jobs"}'), ctx
    )
    assert not outcome.ok and "недоступен" in outcome.output["error"]


# ---------------------------------------------------------------- сквозной сценарий через чат


@pytest.fixture
def provider():
    p = ScriptedProvider()
    app.dependency_overrides[get_llm] = lambda: LLMRouter(
        [Route(p, "gpt-main")], event_timeout=2, total_timeout=10
    )
    yield p
    app.dependency_overrides.pop(get_llm, None)


async def ask(client, conversation_id, text, file_ids=()):
    response = await client.post(
        f"/api/chat/conversations/{conversation_id}/messages",
        json={"text": text, "file_ids": list(file_ids)},
    )
    assert response.status_code == 200, response.text
    return parse_sse(response.text)


async def test_chart_from_exported_data(api, provider, session_factory):
    client, material, supplier = api
    me = (await client.get("/api/auth/me")).json()
    async with session_factory() as s:
        user = await s.get(User, me["id"])
        await _receipt(s, user, supplier, material, day=clock.today() - timedelta(days=5))

    provider.push(
        ToolCallRequested(ToolCall("c1", "export_data", json.dumps({"dataset": "receipts"}))),
        Completed(model="m", finish="tool_calls"),
    )
    chart = provider.add_generated("cntr_1", "/mnt/data/chart.png", PNG)
    provider.add_generated(
        "cntr_1", "/mnt/data/top.csv", f"supplier,kg\nПоставщик-{supplier.id},12.5\n".encode()
    )
    provider.push(
        CodeExecuted(
            code="import pandas as pd\ndf = pd.read_csv('/mnt/data/receipts.csv')\n...",
            container_id="cntr_1",
            outputs=({"type": "logs", "logs": "saved"},),
        ),
        FileCited("cntr_1", chart, "chart.png"),
        TextDelta("Готово: ![график](sandbox:/mnt/data/chart.png) и "),
        TextDelta("[таблица](sandbox:/mnt/data/top.csv)"),
        Completed(model="m"),
    )
    conversation_id = (await client.post("/api/chat/conversations", json={})).json()["id"]

    events = await ask(client, conversation_id, "Построй график поступлений")

    # Второй шаг уже видел выгруженный файл в контейнере
    second = provider.requests[1][1]
    assert second.code_interpreter is not None and second.code_interpreter.file_ids == ("file-1",)
    uploaded_csv = provider.uploaded["file-1"][1].decode()
    assert supplier.name not in uploaded_csv  # поставщик — псевдонимом
    assert provider.deleted == ["file-1"]  # копия у провайдера удалена в конце хода

    code_events = [d for e, d in events if e == "code"]
    assert code_events[0]["code"].startswith("import pandas") and code_events[0]["error"] is False
    files = {d["filename"]: d for e, d in events if e == "file"}
    assert set(files) == {"chart.png", "top.csv"}
    done = events[-1][1]
    assert f"](/api/chat/files/{files['chart.png']['id']})" in done["content"]
    assert "sandbox:" not in done["content"]

    # Картинка отдаётся как есть, в CSV — настоящее имя поставщика
    png = await client.get(files["chart.png"]["url"])
    assert png.content == PNG and png.headers["content-type"] == "image/png"
    csv_text = (await client.get(files["top.csv"]["url"])).content.decode()
    assert csv_text == f"supplier,kg\n{supplier.name},12.5\n"

    # История диалога: код, файлы и итоговый текст сохранены
    message = (await client.get(f"/api/chat/conversations/{conversation_id}")).json()["messages"][1]
    assert message["code_runs"][0]["code"].startswith("import pandas")
    assert {f["filename"] for f in message["files"]} == {"chart.png", "top.csv"}
    assert "sandbox:" not in message["content"]


async def test_failed_code_does_not_break_chat(api, provider):
    client, *_ = api
    provider.push(
        CodeExecuted(
            code="1/0",
            container_id="cntr_err",
            outputs=(
                {
                    "type": "logs",
                    "logs": "Traceback (most recent call last):\n"
                    "ZeroDivisionError: division by zero",
                },
            ),
        ),
        TextDelta("Расчёт не удался: деление на ноль. Уточните данные."),
        Completed(model="m"),
    )
    conversation_id = (await client.post("/api/chat/conversations", json={})).json()["id"]
    events = await ask(client, conversation_id, "Посчитай")
    done = events[-1][1]
    assert done["status"] == "complete"
    assert done["code_runs"][0]["error"] is True
    assert "ZeroDivisionError" in done["code_runs"][0]["logs"]


async def test_lost_container_files_give_a_readable_note(api, provider):
    client, *_ = api

    async def broken(container_id):
        raise RuntimeError("container expired")

    provider.list_generated_files = broken
    provider.push(
        CodeExecuted(code="plot()", container_id="cntr_gone"),
        TextDelta("График готов"),
        Completed(model="m"),
    )
    conversation_id = (await client.post("/api/chat/conversations", json={})).json()["id"]
    done = (await ask(client, conversation_id, "График"))[-1][1]
    assert done["status"] == "complete"
    assert "Не удалось получить файлы" in done["content"]


# ---------------------------------------------------------------- загрузка файла пользователем


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def test_uploaded_excel_is_analysed_with_pseudonyms(api, provider, session_factory):
    client, _, supplier = api
    me = (await client.get("/api/auth/me")).json()
    conversation_id = (await client.post("/api/chat/conversations", json={})).json()["id"]
    upload = await client.post(
        f"/api/chat/conversations/{conversation_id}/files",
        files={
            "file": (
                "продажи.xlsx",
                _xlsx([["Поставщик", "Принял", "Кг"], [supplier.name, me["full_name"], 3]]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    provider.push(TextDelta("Посмотрел файл"), Completed(model="m"))

    await ask(client, conversation_id, "Что в файле?", file_ids=[upload.json()["id"]])

    request = provider.requests[0][1]
    assert request.items[-1].text.endswith(
        "[Файлы пользователя доступны в Code Interpreter: /mnt/data/file-1-продажи.xlsx (Excel)]"
    )
    sent = load_workbook(io.BytesIO(provider.uploaded[request.code_interpreter.file_ids[0]][1]))
    values = [c.value for c in sent.active[2]]
    assert supplier.name not in values and me["full_name"] not in values
    assert values[2] == 3

    message = (await client.get(f"/api/chat/conversations/{conversation_id}")).json()["messages"][0]
    assert [f["filename"] for f in message["files"]] == ["продажи.xlsx"]


@pytest.mark.parametrize(
    ("name", "data", "error"),
    [
        ("virus.exe", b"MZ", "только таблицу CSV или Excel"),
        ("empty.csv", b"", "пустой"),
        ("big.csv", b"a" * (5 * 1024 * 1024 + 1), "больше 5 МБ"),
    ],
    ids=["exe", "empty", "too-big"],
)
async def test_upload_validation(api, name, data, error):
    client, *_ = api
    conversation_id = (await client.post("/api/chat/conversations", json={})).json()["id"]
    response = await client.post(
        f"/api/chat/conversations/{conversation_id}/files", files={"file": (name, data)}
    )
    assert response.status_code == 400
    assert error in response.json()["detail"]


async def test_files_are_private(api, provider, session_factory):
    client, *_ = api
    async with session_factory() as s:
        other = await make_user(s, login=f"ci-{uuid4().hex[:6]}")
        conversation = Conversation(user_id=other.id, title="чужой")
        s.add(conversation)
        await s.commit()
        from app.assistant.chat import save_upload

        foreign = await save_upload(s, other, conversation.id, "секрет.csv", b"a,b\n1,2\n")
    assert (await client.get(f"/api/chat/files/{foreign.id}")).status_code == 404
    mine = (await client.post("/api/chat/conversations", json={})).json()["id"]
    response = await client.post(
        f"/api/chat/conversations/{mine}/messages", json={"text": "x", "file_ids": [foreign.id]}
    )
    assert response.status_code == 404


def test_links_to_files_with_spaces_and_brackets_are_rewritten():
    """M9: «график расхода (1).png» — ссылка на наш файл, а не «файл недоступен»."""
    from types import SimpleNamespace

    from app.assistant.code_interpreter import rewrite_sandbox_links

    chart = SimpleNamespace(id=7, filename="график расхода (1).png")
    text = (
        "[Скачать](sandbox:/mnt/data/график расхода (1).png) и "
        "![](sandbox:/mnt/data/график%20расхода%20%281%29.png)"
    )
    assert rewrite_sandbox_links(text, [chart], None) == (
        "[Скачать](/api/chat/files/7) и ![](/api/chat/files/7)"
    )


@pytest.mark.parametrize(
    ("data", "expected"),
    [
        ("Дата;Сумма\n2026-09-01;10\n2026-09-02;20\n".encode("cp1251"), "разделитель «;»"),
        (b"date,sum\n2026-09-01,10\n2026-09-02,20\n", "разделитель «,»"),
    ],
)
def test_csv_delimiter_is_told_to_the_model(data, expected):
    from app.assistant.code_interpreter import describe_upload

    assert expected in describe_upload("закупки.csv", data)
